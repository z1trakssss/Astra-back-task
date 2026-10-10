from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from .models import MacroWeek, Sale


class SalesInterfaceTests(TestCase):
    """Проверяет API и отчёт администратора."""

    @classmethod
    def setUpTestData(cls):
        """Создаёт две недели и две категории товаров."""
        for day in (7, 14):
            MacroWeek.objects.create(
                week=date(2019, 1, day),
                housing_starts_index=100,
                lumber_price_index=400,
                mortgage_rate=4,
                season_factor=Decimal("0.600"),
            )
        Sale.objects.create(
            source_row=1,
            week_id=date(2019, 1, 7),
            region="Pacific",
            product_category="lumber",
            sku="L-1",
            channel="pro_dealer",
            customer_type="contractor",
            units=2,
            unit_price=Decimal("10.00"),
            revenue=Decimal("20.00"),
        )
        Sale.objects.create(
            source_row=2,
            week_id=date(2019, 1, 14),
            region="Mountain",
            product_category="roofing",
            sku="R-1",
            channel="big_box_diy",
            customer_type="homeowner_diy",
            units=3,
            unit_price=Decimal("5.00"),
            revenue=Decimal("15.00"),
        )

    def test_sales_filter_macro_and_bounds(self):
        """Проверяет фильтр региона и ограничение размера страницы."""
        result = self.client.get("/api/v1/sales/", {"region": "Pacific"}).json()
        self.assertEqual(result["count"], 1)
        self.assertEqual(result["results"][0]["macro"]["season_factor"], "0.600")
        self.assertEqual(
            self.client.get("/api/v1/sales/", {"page_size": 201}).status_code, 400
        )

    def test_aggregate_uses_same_filters(self):
        """Проверяет агрегирование выбранных продаж."""
        result = self.client.get(
            "/api/v1/sales/aggregate/",
            {"region": "Pacific", "group_by": "product_category"},
        ).json()
        self.assertEqual(result["results"][0]["revenue"], "20")
        self.assertEqual(result["results"][0]["units"], 2)
        self.assertEqual(
            self.client.get(
                "/api/v1/sales/aggregate/", {"group_by": "password"}
            ).status_code,
            400,
        )

    def test_export_requires_staff_and_respects_admin_filter(self):
        """Проверяет права доступа и фильтры отчёта."""
        url = "/admin/sales/sale/export/?region__exact=Pacific"
        self.assertEqual(self.client.get(url).status_code, 302)
        user = get_user_model().objects.create_superuser(
            "admin", "a@example.test", "test-password"
        )
        self.client.force_login(user)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        rows = list(
            load_workbook(BytesIO(response.content), read_only=True).active.values
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][:3], ("lumber", 1, 2))

    @override_settings(
        STORAGES={
            "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
            "staticfiles": {
                "BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"
            },
        }
    )
    def test_macro_endpoint_and_admin_dashboard(self):
        """Проверяет недельный ряд и отображение панели администратора."""
        result = self.client.get(
            "/api/v1/macro/",
            {"date_from": "2019-01-07", "date_to": "2019-01-07"},
        ).json()
        self.assertEqual(len(result["results"]), 1)
        user = get_user_model().objects.create_superuser(
            "admin", "a@example.test", "test-password"
        )
        self.client.force_login(user)
        dashboard = self.client.get("/admin/")
        sales = self.client.get("/admin/sales/sale/")
        self.assertEqual(dashboard.status_code, 200)
        self.assertContains(dashboard, "Последние 12 месяцев")
        self.assertContains(sales, "Скачать сводный отчёт XLSX")

    def test_export_respects_date_range_and_search(self):
        """Сверяет XLSX с фильтрами даты и поиском в админке."""
        user = get_user_model().objects.create_superuser(
            "admin", "a@example.test", "test-password"
        )
        self.client.force_login(user)
        response = self.client.get(
            "/admin/sales/sale/export/",
            {
                "week__gte": "2019-01-01",
                "week__lt": "2019-01-10",
                "q": "L-1",
            },
        )
        self.assertEqual(response.status_code, 200)
        rows = list(
            load_workbook(BytesIO(response.content), read_only=True).active.values
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][:3], ("lumber", 1, 2))


class DatasetImportTests(TestCase):
    """Проверяет загрузку исходного архива."""

    def _archive(self, path, revenue, week="2019-01-07"):
        """Создаёт архив с одной неделей и одной продажей."""
        week_number = date.fromisoformat(week).isocalendar().week
        with ZipFile(path, "w") as archive:
            archive.writestr(
                "macro_drivers_weekly.csv",
                "week,housing_starts_index,lumber_price_index,mortgage_rate,season_factor\n"
                f"{week},100,400,4,0.6\n",
            )
            archive.writestr(
                "building_materials_transactions.csv",
                "date,year,month,week_of_year,region,product_category,sku,channel,customer_type,units,unit_price,revenue,housing_starts_index,lumber_price_index,mortgage_rate\n"
                f"{week},2019,1,{week_number},Pacific,lumber,L-1,pro_dealer,contractor,2,10,{revenue},100,400,4\n",
            )

    def test_import_is_repeatable_and_atomic(self):
        """Не создаёт дубли и откатывает архив с неверной суммой."""
        with TemporaryDirectory() as directory:
            source = Path(directory) / "sample.zip"
            self._archive(source, "20.00")
            call_command("import_dataset", source=source, stdout=StringIO())
            call_command("import_dataset", source=source, stdout=StringIO())
            self.assertEqual(Sale.objects.count(), 1)
            self.assertEqual(MacroWeek.objects.count(), 1)
            source_bad = Path(directory) / "bad.zip"
            self._archive(source_bad, "25.00", week="2019-01-14")
            with self.assertRaises(CommandError):
                call_command("import_dataset", source=source_bad, stdout=StringIO())
            self.assertEqual(Sale.objects.count(), 1)
            self.assertEqual(MacroWeek.objects.count(), 1)
